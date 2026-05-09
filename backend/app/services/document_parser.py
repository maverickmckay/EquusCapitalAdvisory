"""
Document parsing service.

Extracts raw text from PDF, DOCX, and plain-text files.
Returns a ParsedDocument with the full text and initial metadata.
"""
from __future__ import annotations
import io
import os
import re
from pathlib import Path
from typing import Optional

from ..models.document import ParsedDocument, DocumentType, UploadedDocument


def parse_document(doc: UploadedDocument, file_path: Path) -> ParsedDocument:
    """Parse a single uploaded file into a ParsedDocument."""
    ext = file_path.suffix.lower()
    full_text = ""
    page_count: Optional[int] = None
    errors = []

    try:
        if ext == ".pdf":
            full_text, page_count = _parse_pdf(file_path)
        elif ext in (".docx", ".doc"):
            full_text = _parse_docx(file_path)
        elif ext in (".txt", ".text", ".md"):
            full_text = file_path.read_text(encoding="utf-8", errors="replace")
        elif ext in (".csv",):
            full_text = file_path.read_text(encoding="utf-8", errors="replace")
        else:
            errors.append(f"Unsupported file type: {ext}. Attempting plain-text read.")
            try:
                full_text = file_path.read_text(encoding="utf-8", errors="replace")
            except Exception as e:
                errors.append(f"Plain-text fallback failed: {e}")
    except Exception as e:
        errors.append(f"Parse error: {e}")

    return ParsedDocument(
        document_id=doc.document_id,
        document_type=doc.document_type,
        full_text=full_text,
        page_count=page_count,
        parse_errors=errors,
    )


def _parse_pdf(path: Path):
    """Try pdfplumber first, then pypdf fallback."""
    text_parts = []
    page_count = None

    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(extracted)
        return "\n\n".join(text_parts), page_count
    except ImportError:
        pass
    except Exception:
        pass

    try:
        import pypdf
        reader = pypdf.PdfReader(str(path))
        page_count = len(reader.pages)
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text_parts.append(extracted)
        return "\n\n".join(text_parts), page_count
    except ImportError:
        pass
    except Exception as e:
        raise RuntimeError(f"PDF parsing failed: {e}")

    return "", page_count


def _parse_docx(path: Path) -> str:
    try:
        import docx
        doc = docx.Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also pull table cell text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n\n".join(paragraphs)
    except ImportError:
        raise RuntimeError("python-docx not installed; cannot parse DOCX files")
    except Exception as e:
        raise RuntimeError(f"DOCX parsing failed: {e}")
